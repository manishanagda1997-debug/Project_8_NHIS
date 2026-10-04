"""Manual-label samples for the taxonomy (decision D-027).

Two samples are drawn from the titles eligible for text analysis:

* **blind test set** — drawn first, as a simple random sample of jobs, so its
  accuracy estimate describes the whole market. It is labelled only after the
  taxonomy, rules, model and threshold are frozen, and scored once.
* **development set** — drawn from the remaining jobs: mostly random, plus a
  few short titles and titles that join two tasks ("logo and website"), so
  hard cases appear while the rules are tuned.

No title key appears twice within a set or in both sets, so the same job
wording can never be in both. Label sheets are Excel files with a category
dropdown; reading them back validates every label.
"""

from __future__ import annotations

import warnings

from pathlib import Path

import pandas as pd

from .taxonomy import CATEGORIES, CATEGORY_IDS

SHEET_COLUMNS = ["sample_id", "row_id", "title", "category", "unsure", "note"]
JOINER = r"\s(?:and|&|\+)\s|/"


class LabelSheetError(ValueError):
    """Raised when a label sheet is missing, malformed or contains invalid labels."""


def _distinct_keys(rows: pd.DataFrame, n: int, seed: int, exclude: set[str]) -> pd.DataFrame:
    """Up to ``n`` random rows with distinct title keys not in ``exclude``.

    Rows are drawn uniformly (a title that is posted often is more likely to
    be drawn, which matches the market), then repeated keys are dropped.
    """
    pool = rows[~rows["title_key"].isin(exclude)]
    shuffled = pool.sample(frac=1.0, random_state=seed)
    return shuffled.drop_duplicates("title_key").head(n)


def draw_samples(frame: pd.DataFrame, rules: dict) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return (blind, development) samples with ``row_id``, ``title`` and ``title_key``."""
    pool = frame.loc[frame["eligible_for_text_analysis"], ["row_id", "title_clean", "title_key", "title_word_count"]]
    pool = pool.rename(columns={"title_clean": "title"}).astype({"title": object, "title_key": object})
    seed = int(rules["seed"])

    blind = _distinct_keys(pool, int(rules["blind_size"]), seed, exclude=set())
    used = set(blind["title_key"])

    short = pool[pool["title_word_count"] <= 2]
    joined = pool[pool["title"].str.contains(JOINER, regex=True)]
    parts = []
    for rows, size, offset in [
        (short, int(rules["dev_short_titles"]), 1),
        (joined, int(rules["dev_joined_titles"]), 2),
        (pool, int(rules["dev_random"]), 3),
    ]:
        part = _distinct_keys(rows, size, seed + offset, exclude=used)
        used |= set(part["title_key"])
        parts.append(part)
    dev = pd.concat(parts).sample(frac=1.0, random_state=seed + 4)  # mix, so hard cases are not grouped

    def finish(sample: pd.DataFrame, prefix: str) -> pd.DataFrame:
        out = sample[["row_id", "title", "title_key"]].reset_index(drop=True)
        out.insert(0, "sample_id", [f"{prefix}{i:03d}" for i in range(1, len(out) + 1)])
        return out

    return finish(blind, "B"), finish(dev, "D")


def write_label_sheet(sample: pd.DataFrame, path: Path) -> Path:
    """Write an Excel label sheet with a category dropdown and a reference tab.

    The sheet shows only the title: no prediction, pay, country or date.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    if path.exists():
        raise FileExistsError(f"{path} already exists; it may contain labels. Move it away first if you really want a new sheet.")
    path.parent.mkdir(parents=True, exist_ok=True)

    book = Workbook()
    sheet = book.active
    sheet.title = "labels"
    sheet.append(SHEET_COLUMNS)
    for row in sample.itertuples(index=False):
        sheet.append([row.sample_id, int(row.row_id), row.title, None, None, None])

    reference = book.create_sheet("categories")
    reference.append(["category", "name", "covers"])
    for category_id, (name, covers) in CATEGORIES.items():
        reference.append([category_id, name, covers])

    last = len(sample) + 1
    category_list = DataValidation(type="list", formula1=f"=categories!$A$2:$A${len(CATEGORIES) + 1}", allow_blank=True)
    category_list.error = "Choose a category from the list (see the 'categories' tab)."
    unsure_list = DataValidation(type="list", formula1='"yes"', allow_blank=True)
    sheet.add_data_validation(category_list)
    sheet.add_data_validation(unsure_list)
    category_list.add(f"D2:D{last}")
    unsure_list.add(f"E2:E{last}")

    header_fill = PatternFill("solid", fgColor="DDE7F5")
    for ws, widths in [(sheet, [10, 9, 80, 30, 9, 40]), (reference, [30, 42, 100])]:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for index, width in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + index)].width = width
        ws.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    book.save(path)
    return path


def read_label_sheet(path: Path, require_complete: bool = True) -> pd.DataFrame:
    """Read a label sheet and validate it.

    Every category must be one of the taxonomy IDs; ``unsure`` becomes a
    boolean. With ``require_complete``, every row must have a category.
    """
    path = Path(path)
    if not path.is_file():
        raise LabelSheetError(f"label sheet not found: {path}")
    with warnings.catch_warnings():
        # Excel stores a dropdown that points to another tab as an "extension"
        # that openpyxl cannot read; the values themselves are read and checked below.
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        table = pd.read_excel(path, sheet_name="labels", dtype={"sample_id": str, "title": str, "category": str, "note": str})
    table = table.dropna(how="all")      # formatted but empty rows below the sample
    missing = [c for c in SHEET_COLUMNS if c not in table.columns]
    if missing:
        raise LabelSheetError(f"{path.name} lacks columns {missing}")
    table["category"] = table["category"].str.strip()
    unknown = sorted(set(table["category"].dropna()) - set(CATEGORY_IDS))
    if unknown:
        raise LabelSheetError(f"{path.name} has unknown categories: {unknown}")
    if require_complete and table["category"].isna().any():
        empty = table.loc[table["category"].isna(), "sample_id"].tolist()
        raise LabelSheetError(f"{path.name} has {len(empty)} unlabelled rows, e.g. {empty[:5]}")
    if table["sample_id"].duplicated().any():
        raise LabelSheetError(f"{path.name} has duplicate sample IDs")
    table["unsure"] = table["unsure"].astype(str).str.strip().str.lower().eq("yes")
    table["row_id"] = table["row_id"].astype("int64")
    return table[SHEET_COLUMNS]


# ------------------------------------------------- recommendation relevance

RELEVANCE_COLUMNS = ["item_id", "profile_id", "title", "relevance", "note", "title_key"]
RELEVANCE_GUIDE = (
    ("2", "Relevant", "This freelancer could do the main work of this job with the skills listed."),
    ("1", "Partly relevant", "The job uses some of their skills, but the main work needs a different skill, "
                             "or the title is too vague to be sure."),
    ("0", "Not relevant", "The job does not fit this freelancer's skills."),
)


def write_relevance_sheet(items: pd.DataFrame, profiles: list[dict], path: Path) -> Path:
    """Write the blinded relevance sheet (decision D-034).

    One row per (profile, title key) found by any method, shuffled within each
    profile. The sheet shows only the profile and the title: no method, rank,
    score, pay, country or date. ``title_key`` is a hidden matching column.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    if path.exists():
        raise FileExistsError(f"{path} already exists; it may contain judgements. Move it away first if you really want a new sheet.")
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "relevance"
    sheet.append(RELEVANCE_COLUMNS)
    for row in items.itertuples(index=False):
        sheet.append([row.item_id, row.profile_id, row.title, None, None, row.title_key])

    guide = book.create_sheet("guide")
    guide.append(["profile_id", "profile", "skills typed in"])
    for profile in profiles:
        guide.append([profile["id"], profile["name"], profile["skills"]])
    guide.append([])
    guide.append(["relevance", "meaning", "rule"])
    for entry in RELEVANCE_GUIDE:
        guide.append(list(entry))
    guide.append([])
    guide.append(["note", "Judge the title only (title-level relevance), against the profile's skills. "
                          "The same title can appear under several profiles; judge it for each profile separately."])

    last = len(items) + 1
    choice = DataValidation(type="list", formula1='"2,1,0"', allow_blank=True)
    choice.error = "Choose 2 (relevant), 1 (partly relevant) or 0 (not relevant)."
    sheet.add_data_validation(choice)
    choice.add(f"D2:D{last}")
    header_fill = PatternFill("solid", fgColor="DDE7F5")
    for ws, widths in [(sheet, [9, 10, 80, 11, 40, 30]), (guide, [12, 28, 100])]:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for index, width in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + index)].width = width
        ws.freeze_panes = "A2"
    sheet.column_dimensions["F"].hidden = True
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    book.save(path)
    return path


def read_relevance_sheet(path: Path, require_complete: bool = True) -> pd.DataFrame:
    """Read and validate the relevance sheet: every judgement must be 0, 1 or 2."""
    path = Path(path)
    if not path.is_file():
        raise LabelSheetError(f"relevance sheet not found: {path}")
    table = pd.read_excel(path, sheet_name="relevance", dtype={"item_id": str, "profile_id": str, "title": str,
                                                                 "note": str, "title_key": str})
    table = table.dropna(how="all")
    missing = [c for c in RELEVANCE_COLUMNS if c not in table.columns]
    if missing:
        raise LabelSheetError(f"{path.name} lacks columns {missing}")
    values = pd.to_numeric(table["relevance"], errors="coerce")
    bad = table.loc[table["relevance"].notna() & ~values.isin([0, 1, 2]), "item_id"].tolist()
    if bad:
        raise LabelSheetError(f"{path.name} has invalid relevance values (use 0, 1 or 2), e.g. {bad[:5]}")
    if require_complete and values.isna().any():
        empty = table.loc[values.isna(), "item_id"].tolist()
        raise LabelSheetError(f"{path.name} has {len(empty)} unjudged rows, e.g. {empty[:5]}")
    if table.duplicated(["profile_id", "title_key"]).any() or table["item_id"].duplicated().any():
        raise LabelSheetError(f"{path.name} has duplicate items")
    table["relevance"] = values
    return table[RELEVANCE_COLUMNS]
