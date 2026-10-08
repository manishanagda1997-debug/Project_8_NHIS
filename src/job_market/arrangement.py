"""Work-arrangement and time-zone wording audit: frame, split and development sample (decision D-041).

This module covers only the **development** stage. ``title_frame`` holds
every title key in the window (both sides, in memory, to assign the split);
``development_sample`` returns and the pilot sheet shows development titles
only, so rule development never sees an evaluation-side title:

1. ``title_frame``: one row per distinct ``title_key`` in the time-series
   window, with its posting count and the most frequent cleaned title.
2. ``assign_side``: every title key goes to development or evaluation once,
   by a seeded SHA-256 of the key. The result does not depend on row order
   or on which other titles exist, so a key can never be on both sides.
3. ``development_sample``: mutually exclusive strata on the development side
   (priority order D1 > D2 > ... > D6); D1, the hybrid audit, is taken in
   full, the others by seeded simple random sampling.

The evaluation sample is designed only after the development rules are
frozen (see the decision record); nothing here draws it.
"""

from __future__ import annotations

import hashlib
import re
from typing import Any

import numpy as np
import pandas as pd

from .trends import signal_flags

STRATA = (
    ("D1", "hybrid audit: old hybrid pattern (all development titles)"),
    ("D2", "old time-zone match"),
    ("D3", "old on-site / in person / local match"),
    ("D4", "old remote or work-from-home match"),
    ("D5", "alternative-wording candidate, no old match"),
    ("D6", "no old match and no candidate"),
)
SHEET_COLUMNS = ["sample_id", "title", "remote_or_home", "on_site_required", "time_zone_required", "note"]
GUIDE = "docs/arrangement_labelling_guide.md"
SIZE_KEYS = {"D2": "dev_time_zone", "D3": "dev_on_site", "D4": "dev_remote_home", "D5": "dev_candidates",
             "D6": "dev_unmatched"}


def title_frame(jobs: pd.DataFrame) -> pd.DataFrame:
    """Distinct title keys in the time-series window, deduplicated."""
    window = jobs.loc[jobs["eligible_for_timeseries"] & jobs["title_key"].fillna("").str.len().gt(0),
                      ["row_id", "title_key", "title_clean"]]
    postings = window.groupby("title_key").size().rename("postings")
    # The title shown to the labeller: the most frequent cleaned title of the
    # key; ties go to the earliest row, so the choice is reproducible.
    counts = (window.groupby(["title_key", "title_clean"]).agg(n=("row_id", "size"), first_row=("row_id", "min"))
              .reset_index().sort_values(["title_key", "n", "first_row"], ascending=[True, False, True]))
    shown = counts.drop_duplicates("title_key").set_index("title_key")["title_clean"].rename("title")
    return pd.concat([shown, postings], axis=1).reset_index().rename(columns={"index": "title_key"})


def assign_side(title_keys: pd.Series, seed: int, development_share: float) -> pd.Series:
    """'development' or 'evaluation' for each key, from a seeded hash of the key alone."""
    def side(key: str) -> str:
        value = int.from_bytes(hashlib.sha256(f"{seed}:{key}".encode("utf-8")).digest()[:8], "big") / 2 ** 64
        return "development" if value < development_share else "evaluation"
    return title_keys.map(side)


def candidate_mask(title_keys: pd.Series, patterns: list[str]) -> pd.Series:
    text = " " + title_keys.fillna("").astype(str) + " "
    combined = "|".join(f"(?:{p})" for p in patterns)
    return text.str.contains(rf"\s(?:{combined})\s", regex=True)


def strata(frame: pd.DataFrame, rules: dict[str, Any]) -> pd.Series:
    """One stratum per title key, by fixed priority (mutually exclusive)."""
    flags = signal_flags(frame["title_key"])
    text = " " + frame["title_key"].astype(str) + " "
    hybrid = text.str.contains(rf"\s(?:{rules['hybrid_audit_pattern']})\s", regex=True)
    remote_home = flags["remote"] | flags["work from home"]
    candidate = candidate_mask(frame["title_key"], list(rules["candidate_patterns"]))
    choices = [hybrid, flags["time-zone requirement"], flags["on-site / in person / local"], remote_home, candidate]
    return pd.Series(np.select(choices, ["D1", "D2", "D3", "D4", "D5"], default="D6"), index=frame.index)


def development_sample(frame: pd.DataFrame, rules: dict[str, Any]) -> tuple[pd.DataFrame, pd.DataFrame]:
    """(sample, stratum summary) for the development side only.

    The sample is shuffled with the seed and numbered; the first
    ``pilot_titles`` rows form the pilot. Columns that reveal the stratum or
    the selection reason are kept for the record, never shown to the labeller.
    """
    seed = int(rules["seed"])
    frame = frame.assign(side=assign_side(frame["title_key"], seed, float(rules["development_share"])))
    dev = frame[frame["side"] == "development"].copy()
    dev["stratum"] = strata(dev, rules)
    rng = np.random.default_rng(seed)
    parts, summary = [], []
    for code, reason in STRATA:
        pool = dev[dev["stratum"] == code].sort_values("title_key")      # sorted: reproducible draws
        if code == "D1":
            chosen = pool
        else:
            n = min(int(rules[SIZE_KEYS[code]]), len(pool))
            chosen = pool.iloc[np.sort(rng.choice(len(pool), size=n, replace=False))] if n else pool.iloc[:0]
        parts.append(chosen.assign(selection_reason=reason))
        summary.append({"stratum": code, "description": reason, "development_titles": len(pool),
                        "sampled": len(chosen), "taken_in_full": code == "D1" or len(chosen) == len(pool)})
    sample = pd.concat(parts, ignore_index=True)
    sample = sample.iloc[rng.permutation(len(sample))].reset_index(drop=True)
    sample.insert(0, "sample_id", [f"A{i + 1:03d}" for i in range(len(sample))])
    sample["pilot"] = sample.index < int(rules["pilot_titles"])
    return sample[["sample_id", "title_key", "title", "postings", "stratum", "selection_reason", "pilot"]], \
        pd.DataFrame(summary)


def write_arrangement_sheet(sample: pd.DataFrame, path) -> "Path":
    """Excel sheet for blind labelling: sample ID and title only, three yes/no/unclear dropdowns.

    No stratum, rule match, count or date is written. Refuses to overwrite.
    """
    from pathlib import Path
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    if path.exists():
        raise FileExistsError(f"{path} already exists; it may contain labels")
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "labels"
    sheet.append(SHEET_COLUMNS)
    for row in sample.itertuples(index=False):
        sheet.append([row.sample_id, row.title, None, None, None, None])
    choice = DataValidation(type="list", formula1='"yes,no,unclear"', allow_blank=True)
    choice.error = "Choose yes, no or unclear (see docs/arrangement_labelling_guide.md)."
    sheet.add_data_validation(choice)
    choice.add(f"C2:E{len(sample) + 1}")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE7F5")
    for column, width in zip("ABCDEF", (10, 80, 16, 16, 16, 40)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    book.save(path)
    return path


def selection_record(sample: pd.DataFrame, summary: pd.DataFrame, rules: dict[str, Any], source_sha256: str,
                     jobs_sha256: str, guide_sha256: str) -> dict[str, Any]:
    """What was drawn, from what, and how: written next to the sample."""
    frozen = {key: rules[key] for key in sorted(rules)}
    return {"decision": "D-041", "stage": "development", "source_sha256": source_sha256,
            "jobs_clean_sha256": jobs_sha256, "guide_sha256": guide_sha256, "seed": int(rules["seed"]),
            "development_share": float(rules["development_share"]),
            "settings_sha256": hashlib.sha256(repr(frozen).encode("utf-8")).hexdigest(),
            "titles": int(len(sample)), "pilot_titles": int(sample["pilot"].sum()),
            "strata": summary.to_dict("records")}
