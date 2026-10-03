"""Label samples: sizes, no shared title keys, determinism, and safe sheets."""

import pandas as pd
import pytest

from job_market.config import load_settings
from job_market.labels import LabelSheetError, draw_samples, read_label_sheet, write_label_sheet
from job_market.taxonomy import CATEGORY_IDS, OTHER, TAXONOMY_VERSION

RULES = {"seed": 7, "blind_size": 30, "dev_random": 15, "dev_short_titles": 5, "dev_joined_titles": 5}


def fake_jobs(n=400):
    titles = []
    for i in range(n):
        if i % 10 == 0:
            titles.append(f"logo {i}")                      # short (2 words)
        elif i % 10 == 1:
            titles.append(f"logo and website number {i}")   # joined
        elif i % 10 == 2:
            titles.append("video editor for youtube")       # one repeated title key
        else:
            titles.append(f"data analyst for project {i}")
    return pd.DataFrame({
        "row_id": range(n),
        "title_clean": titles,
        "title_key": [t.lower() for t in titles],
        "title_word_count": [len(t.split()) for t in titles],
        "eligible_for_text_analysis": [True] * (n - 5) + [False] * 5,
    })


def test_sizes_and_ids():
    blind, dev = draw_samples(fake_jobs(), RULES)
    assert len(blind) == 30 and len(dev) == 25
    assert blind["sample_id"].iloc[0] == "B001" and dev["sample_id"].str.startswith("D").all()


def test_no_title_key_repeats_within_or_across_sets():
    blind, dev = draw_samples(fake_jobs(), RULES)
    assert blind["title_key"].is_unique and dev["title_key"].is_unique
    assert not set(blind["title_key"]) & set(dev["title_key"])


def test_development_set_contains_hard_cases():
    _, dev = draw_samples(fake_jobs(), RULES)
    assert (dev["title"].str.split().str.len() <= 2).sum() >= 5
    assert dev["title"].str.contains(" and ").sum() >= 5


def test_only_eligible_rows_are_sampled():
    jobs = fake_jobs()
    blind, dev = draw_samples(jobs, RULES)
    ineligible = set(jobs.loc[~jobs["eligible_for_text_analysis"], "row_id"])
    assert not ineligible & (set(blind["row_id"]) | set(dev["row_id"]))


def test_same_seed_same_samples():
    first = draw_samples(fake_jobs(), RULES)
    second = draw_samples(fake_jobs(), RULES)
    pd.testing.assert_frame_equal(first[0], second[0])
    pd.testing.assert_frame_equal(first[1], second[1])


def test_config_and_code_agree_on_taxonomy_version():
    from conftest import REPO_ROOT

    assert load_settings(REPO_ROOT).taxonomy_version == TAXONOMY_VERSION
    assert len(CATEGORY_IDS) == 23 and OTHER in CATEGORY_IDS


# ------------------------------------------------------------------ sheets


def labelled_copy(path, values):
    from openpyxl import load_workbook

    book = load_workbook(path)
    sheet = book["labels"]
    for row, (category, unsure) in zip(range(2, sheet.max_row + 1), values):
        sheet.cell(row=row, column=4, value=category)
        sheet.cell(row=row, column=5, value=unsure)
    book.save(path)


def test_sheet_round_trip(tmp_path):
    pytest.importorskip("openpyxl")
    _, dev = draw_samples(fake_jobs(), RULES)
    path = write_label_sheet(dev.head(3), tmp_path / "dev.xlsx")
    labelled_copy(path, [("design_creative", None), ("web_development", "yes"), (OTHER, None)])
    table = read_label_sheet(path)
    assert table["category"].tolist() == ["design_creative", "web_development", OTHER]
    assert table["unsure"].tolist() == [False, True, False]
    assert table["title"].tolist() == dev.head(3)["title"].tolist()


def test_existing_sheet_is_never_overwritten(tmp_path):
    pytest.importorskip("openpyxl")
    _, dev = draw_samples(fake_jobs(), RULES)
    path = write_label_sheet(dev.head(2), tmp_path / "dev.xlsx")
    with pytest.raises(FileExistsError):
        write_label_sheet(dev.head(2), path)


def test_unknown_or_missing_labels_are_rejected(tmp_path):
    pytest.importorskip("openpyxl")
    _, dev = draw_samples(fake_jobs(), RULES)
    path = write_label_sheet(dev.head(2), tmp_path / "dev.xlsx")
    labelled_copy(path, [("designs", None), ("web_development", None)])
    with pytest.raises(LabelSheetError, match="unknown categories"):
        read_label_sheet(path)
    path2 = write_label_sheet(dev.head(2), tmp_path / "dev2.xlsx")
    labelled_copy(path2, [("legal", None), (None, None)])
    with pytest.raises(LabelSheetError, match="unlabelled"):
        read_label_sheet(path2)
    assert len(read_label_sheet(path2, require_complete=False)) == 2
