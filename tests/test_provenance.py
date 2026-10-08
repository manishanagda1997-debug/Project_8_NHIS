"""Label provenance and the recommender serving approval (decisions D-039, D-040)."""

from __future__ import annotations

import pandas as pd
import pytest

from job_market import provenance as pv
from job_market.manifest import sha256_file
from job_market.steps.recommend import RELEVANCE_SHEET, current_evaluation, decide_serving, historical_guard

SOURCE = "a" * 64
OTHER_SOURCE = "b" * 64


# ------------------------------------------------------------ label status

def sheet(folder, name="blind_labels.xlsx", content=b"labels"):
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_bytes(content)
    return name


def test_missing_sheet(tmp_path):
    assert pv.label_status(tmp_path, "blind_labels.xlsx", SOURCE)[0] == pv.MISSING


def test_missing_or_invalid_source_metadata_is_never_scored(tmp_path):
    name = sheet(tmp_path)
    assert pv.label_status(tmp_path, name, SOURCE)[0] == pv.NO_IDENTITY             # no label_sources.toml
    (tmp_path / pv.LABEL_SOURCES_FILE).write_text("[other.xlsx]\nsheet_sha256 = \"x\"\n", encoding="utf-8")
    assert pv.label_status(tmp_path, name, SOURCE)[0] == pv.NO_IDENTITY             # no entry for this sheet
    (tmp_path / pv.LABEL_SOURCES_FILE).write_text(
        f'["{name}"]\nsheet_sha256 = "not-a-hash"\nsource_sha256 = "{SOURCE}"\n', encoding="utf-8")
    assert pv.label_status(tmp_path, name, SOURCE)[0] == pv.NO_IDENTITY             # malformed hash
    (tmp_path / pv.LABEL_SOURCES_FILE).write_text("this is not [toml", encoding="utf-8")
    assert pv.label_status(tmp_path, name, SOURCE)[0] == pv.NO_IDENTITY             # unreadable file


def test_changed_sheet_is_malformed_and_other_source_is_rejected(tmp_path):
    name = sheet(tmp_path)
    pv.bind_label_sheet(tmp_path, name, SOURCE)
    assert pv.label_status(tmp_path, name, SOURCE)[0] == pv.SCORED
    assert pv.label_status(tmp_path, name, OTHER_SOURCE)[0] == pv.WRONG_SOURCE
    (tmp_path / name).write_bytes(b"edited labels")
    assert pv.label_status(tmp_path, name, SOURCE)[0] == pv.MALFORMED


def relevance_sheet(folder, pairs):
    from job_market.labels import write_relevance_sheet
    from openpyxl import load_workbook
    items = pd.DataFrame({"item_id": [f"R{i:03d}" for i in range(len(pairs))], "profile_id": [p for p, _ in pairs],
                          "title": [t for _, t in pairs], "title_key": [t for _, t in pairs]})
    path = write_relevance_sheet(items, [{"id": "P01", "name": "x", "skills": "a"}], folder / RELEVANCE_SHEET)
    book = load_workbook(path)
    for row in range(2, len(pairs) + 2):
        book["relevance"][f"D{row}"] = 2
    book.save(path)


RESULTS = pd.DataFrame({"profile_id": ["P01", "P01"], "method": ["keyword", "keyword"], "rank": [1, 2],
                        "row_id": [10, 11], "title_key": ["logo design", "brand identity"],
                        "title": ["Logo design", "Brand identity"]})
RULES = {"eval_top_k": 5, "profiles": [{"id": "P01", "name": "x"}], "methods": ["keyword"], "tie_threshold": 0.02}


def test_other_source_blocks_scoring_even_when_every_pair_is_judged(tmp_path):
    relevance_sheet(tmp_path, [("P01", "logo design"), ("P01", "brand identity")])
    pv.bind_label_sheet(tmp_path, RELEVANCE_SHEET, SOURCE)
    called = []

    def results():
        called.append(True)
        return RESULTS

    status, _, scored = current_evaluation(tmp_path, OTHER_SOURCE, results, RULES)
    assert status == pv.WRONG_SOURCE and scored is None and not called           # never even retrieved
    status, _, scored = current_evaluation(tmp_path, SOURCE, results, RULES)
    assert status == pv.SCORED and scored is not None and scored[3] == "keyword"


# ------------------------------------------------------------ build-level label checks

def build(settings):
    from job_market.pipeline import run_pipeline
    return run_pipeline(settings, log=lambda _: None)


def dev_sheet(settings, titles=("Logo designer", "Virtual assistant")):
    from job_market.labels import write_label_sheet
    from openpyxl import load_workbook
    folder = settings.path("labels_dir")
    folder.mkdir(parents=True, exist_ok=True)
    path = write_label_sheet(pd.DataFrame({"sample_id": ["D001", "D002"], "row_id": [1, 3], "title": list(titles)}),
                             folder / "dev_labels.xlsx")
    book = load_workbook(path)
    book["labels"]["D2"], book["labels"]["D3"] = "design_creative", "admin_virtual_assistance"
    book.save(path)
    return folder


def reorder_raw(settings):
    raw = settings.path("raw_csv")
    lines = raw.read_text(encoding="utf-8").splitlines()
    raw.write_text("\n".join([lines[0], lines[2], lines[1], *lines[3:]]) + "\n", encoding="utf-8")


def evaluate_metrics(manifest):
    return {step["name"]: step["metrics"] for step in manifest["steps"]}["evaluate"]


def test_reordered_extract_of_the_same_length_is_rejected(mini_project):
    folder = dev_sheet(mini_project)
    pv.bind_label_sheet(folder, "dev_labels.xlsx", sha256_file(mini_project.path("raw_csv")))
    assert evaluate_metrics(build(mini_project))["dev_labels"] == "scored"
    reorder_raw(mini_project)
    assert evaluate_metrics(build(mini_project))["dev_labels"].startswith(pv.WRONG_SOURCE)


def test_matching_source_with_a_changed_title_stops_the_build(mini_project):
    folder = dev_sheet(mini_project, titles=("Logo designer", "Bookkeeper"))
    pv.bind_label_sheet(folder, "dev_labels.xlsx", sha256_file(mini_project.path("raw_csv")))
    with pytest.raises(pv.ProvenanceError, match="no longer match"):
        build(mini_project)


def test_rebinding_to_a_reordered_extract_is_caught_by_the_title_check(mini_project):
    folder = dev_sheet(mini_project)
    reorder_raw(mini_project)                                          # rows 1 and 2 swap places ...
    pv.bind_label_sheet(folder, "dev_labels.xlsx", sha256_file(mini_project.path("raw_csv")))   # ... record edited
    with pytest.raises(pv.ProvenanceError):
        build(mini_project)


# ------------------------------------------------------------ historical rankings

def write_historical(folder, items, source):
    folder.mkdir(parents=True, exist_ok=True)
    items[pv.RANKING_FIELDS].to_csv(folder / "hist.csv", index=False)
    (folder / pv.HISTORICAL_RANKINGS_FILE).write_text(
        f'file = "hist.csv"\nsha256 = "{sha256_file(folder / "hist.csv")}"\nsource_sha256 = "{source}"\n'
        'build_id = "H1"\n', encoding="utf-8")


def test_historical_ranking_identity_only_on_the_historical_extract(tmp_path):
    write_historical(tmp_path, RESULTS, SOURCE)
    assert historical_guard(tmp_path, SOURCE, RESULTS, SOURCE) == "identical to build H1"
    changed = RESULTS.assign(row_id=[10, 99])
    with pytest.raises(pv.ProvenanceError, match="differ"):
        historical_guard(tmp_path, SOURCE, changed, SOURCE)
    assert historical_guard(tmp_path, OTHER_SOURCE, changed, SOURCE).startswith("not applicable")


def test_missing_historical_record_on_the_historical_extract_is_an_error(tmp_path):
    with pytest.raises(pv.ProvenanceError, match="record .* is missing"):
        historical_guard(tmp_path, SOURCE, RESULTS, SOURCE)            # deleting the record cannot skip the check
    assert historical_guard(tmp_path, OTHER_SOURCE, RESULTS, SOURCE).startswith("not applicable")
    write_historical(tmp_path, RESULTS, OTHER_SOURCE)                  # a record for another extract
    with pytest.raises(pv.ProvenanceError, match="another extract"):
        historical_guard(tmp_path, SOURCE, RESULTS, SOURCE)


# ------------------------------------------------------------ serving decision

EVIDENCE = {"build_id": "20261005T122621Z-90e95d23f3", "source_sha256": SOURCE, "labels_sha256": "c" * 64,
            "rankings_file": "recommender_rankings_20261005T122621Z.csv", "rankings_sha256": "d" * 64,
            "selection_rule": "highest mean nDCG@5; ties below 0.02 by strict Precision@5; then simpler"}


def manifest_with(metrics):
    return {"run_id": "R1", "steps": [{"name": "recommend", "metrics": metrics}]}


def approved(settings, method="tfidf_hybrid"):
    pv.write_approval(settings, method, EVIDENCE)
    return pv.load_approval(settings)


def build_metrics(settings, winner, label_state="scored"):
    approval = None
    try:
        approval = pv.load_approval(settings)
        problem = None if approval is None else "; ".join(pv.scope_problems(approval, settings)) or None
    except pv.ApprovalError as error:
        problem = f"invalid: {error}"
    return decide_serving(winner, label_state, approval, problem)


def test_build_with_approval_pending_cannot_be_served(mini_project):
    metrics = build_metrics(mini_project, "tfidf_hybrid")
    assert metrics["recommend_approval_status"] == "pending" and metrics["recommend_served_method"] == "tfidf_hybrid"
    assert "pending" not in metrics["recommend_method_status"]
    with pytest.raises(pv.ServingError, match="approval pending"):
        pv.serving_decision(mini_project, manifest_with(metrics))


def test_post_build_approval_of_the_recorded_method_serves_without_rebuild(mini_project):
    metrics = build_metrics(mini_project, "tfidf_hybrid")
    approved(mini_project)
    decision = pv.serving_decision(mini_project, manifest_with(metrics))
    assert decision["method"] == "tfidf_hybrid" and decision["approval_note"] == "approval recorded after this build"
    assert decision["build_approval_ref"] is None and decision["approval_ref"]


def test_changed_approval_cannot_switch_the_method_under_existing_outputs(mini_project):
    metrics = build_metrics(mini_project, "tfidf_hybrid")
    approved(mini_project, "keyword")
    with pytest.raises(pv.ServingError, match="approval changed after this build"):
        pv.serving_decision(mini_project, manifest_with(metrics))


def test_different_source_build_serves_the_approved_method(mini_project):
    approved(mini_project)
    metrics = build_metrics(mini_project, None, label_state=pv.WRONG_SOURCE)
    assert metrics["recommend_method_status"].startswith("chosen by the earlier relevance test")
    decision = pv.serving_decision(mini_project, manifest_with(metrics))   # no historical source match needed
    assert decision["method"] == "tfidf_hybrid" and "no new relevance test" in decision["method_status"]


def test_winner_mismatch_is_refused(mini_project):
    approved(mini_project)
    metrics = build_metrics(mini_project, "keyword")
    assert metrics["recommend_winner_mismatch"]
    with pytest.raises(pv.ServingError, match="a person must decide"):
        pv.serving_decision(mini_project, manifest_with(metrics))


def test_invalid_approval_is_never_pending(mini_project):
    path = mini_project.root / pv.APPROVAL_FILE
    path.write_text("approved_method = [broken", encoding="utf-8")
    metrics = build_metrics(mini_project, "tfidf_hybrid")
    assert metrics["recommend_approval_status"].startswith("invalid")
    with pytest.raises(pv.ServingError, match="approval invalid"):
        pv.serving_decision(mini_project, manifest_with({**metrics}))
    approved(mini_project, "tfidf_hybrid")
    path.write_text(path.read_text(encoding="utf-8").replace('"tfidf_hybrid"', '"magic"'), encoding="utf-8")
    with pytest.raises(pv.ServingError, match="not one of the configured methods"):
        pv.serving_decision(mini_project, manifest_with(build_metrics(mini_project, None)))


def test_changed_code_or_settings_stop_serving(mini_project):
    metrics = build_metrics(mini_project, "tfidf_hybrid")
    approved(mini_project)
    text = mini_project.root / "src" / "job_market" / "text.py"
    original = text.read_text(encoding="utf-8")
    text.write_text(original + "\n# changed\n", encoding="utf-8")
    with pytest.raises(pv.ServingError, match="code changed"):
        pv.serving_decision(mini_project, manifest_with(metrics))
    text.write_text(original, encoding="utf-8")
    config = mini_project.root / "config" / "pipeline.toml"
    config.write_text(config.read_text(encoding="utf-8").replace("results = 10", "results = 11"), encoding="utf-8")
    from job_market.config import load_settings
    with pytest.raises(pv.ServingError, match="settings changed"):
        pv.serving_decision(load_settings(mini_project.root), manifest_with(metrics))


def test_different_library_versions_warn_normally_and_stop_in_strict_mode(mini_project, monkeypatch):
    metrics = build_metrics(mini_project, "tfidf_hybrid")
    approved(mini_project)
    running = {**pv.library_versions(), "numpy": "0.0"}
    monkeypatch.setattr(pv, "library_versions", lambda: running)
    decision = pv.serving_decision(mini_project, manifest_with(metrics))           # normal mode: serves
    assert "warning: library versions differ" in decision["approval_note"] and "numpy" in decision["approval_note"]
    monkeypatch.setenv("JOB_MARKET_STRICT", "1")
    with pytest.raises(pv.ServingError, match="strict mode"):
        pv.serving_decision(mini_project, manifest_with(metrics))


def test_no_method_to_serve_is_refused(mini_project):
    metrics = build_metrics(mini_project, None, label_state=pv.MISSING)
    assert metrics["recommend_served_method"] == "none"
    with pytest.raises(pv.ServingError, match="no method to serve"):
        pv.serving_decision(mini_project, manifest_with(metrics))


def test_app_refuses_a_stale_build(mini_project):
    from job_market_web.build import AppBuildError, load_service
    build(mini_project)
    config = mini_project.root / "config" / "pipeline.toml"
    config.write_text(config.read_text(encoding="utf-8") + "\n# edited\n", encoding="utf-8")
    from job_market.config import load_settings
    with pytest.raises(AppBuildError, match="stale"):
        load_service(load_settings(mini_project.root))


@pytest.mark.parametrize("missing", ["build_id", "source_sha256", "labels_sha256", "rankings_file", "rankings_sha256",
                                     "selection_rule"])
def test_incomplete_or_invalid_evidence_is_rejected(mini_project, missing):
    pv.write_approval(mini_project, "tfidf_hybrid", {k: v for k, v in EVIDENCE.items() if k != missing})
    with pytest.raises(pv.ApprovalError, match=missing):
        pv.load_approval(mini_project)


def test_short_hash_in_evidence_is_rejected(mini_project):
    pv.write_approval(mini_project, "tfidf_hybrid", {**EVIDENCE, "labels_sha256": "90c9de4a"})
    with pytest.raises(pv.ApprovalError, match="labels_sha256"):
        pv.load_approval(mini_project)


# ------------------------------------------------------------ real loader + API

def served_project(settings, metrics_extra=None, method="tfidf_hybrid"):
    """A processed folder with the files the app reads, recorded in a manifest
    whose fingerprints match the project, plus recommend metrics."""
    import json
    from test_web import jobs_table
    from job_market.fingerprint import project_code_sha256
    from job_market.manifest import describe_output
    processed = settings.path("processed_dir")
    processed.mkdir(parents=True, exist_ok=True)
    jobs_table().to_parquet(processed / "recommend_jobs.parquet", index=False)
    pd.DataFrame({"keyword": ["power bi"], "country_category_pct": [8.0],
                  "country_category_p_fdr": [0.01]}).to_csv(processed / "pay_keywords_hourly.csv", index=False)
    pd.DataFrame(columns=["keyword", "change_pp", "relative_change_pct", "p_fdr"]).to_csv(
        processed / "trend_keywords.csv", index=False)
    outputs = [describe_output(processed / name, settings.root) for name in
               ("recommend_jobs.parquet", "pay_keywords_hourly.csv", "trend_keywords.csv")]
    metrics = build_metrics(settings, method)
    metrics.update(metrics_extra or {})
    from job_market.manifest import MANIFEST_VERSION
    manifest = {"manifest_version": MANIFEST_VERSION, "run_id": "R-int", "fingerprints": {"code_sha256": project_code_sha256(settings),
                                                    "config_sha256": settings.config_sha256},
                "outputs": outputs, "steps": [{"name": "recommend", "metrics": metrics}]}
    settings.path("manifest").write_text(json.dumps(manifest), encoding="utf-8")
    return processed


def test_real_loader_serves_an_approved_build_with_the_same_status_everywhere(mini_project):
    from job_market_web import create_app
    from job_market_web.build import load_service
    served_project(mini_project)
    approved(mini_project)
    service = load_service(mini_project)
    client = create_app(service=service).test_client()
    health = client.get("/api/health").get_json()
    rec = client.get("/api/recommend?skills=logo design").get_json()
    row = int(service.index.jobs["row_id"].iloc[5])
    similar = client.get(f"/api/similar/{row}").get_json()
    fields = ("method", "method_status", "approval_note", "approval_ref", "build_approval_ref")
    assert health["method_status"] == "chosen by the relevance test on this data"      # no stale "pending"
    assert health["approval_note"] == "approval recorded after this build"
    for payload in (rec, similar):
        assert {k: payload[k] for k in fields} == {k: health[k] for k in fields}
    page = client.get("/").get_data(as_text=True)
    assert health["method_status"] in page


def test_real_loader_refuses_pending_approval_and_a_corrupted_output(mini_project):
    from job_market_web.build import AppBuildError, load_service
    processed = served_project(mini_project)
    with pytest.raises(AppBuildError, match="approval pending"):
        load_service(mini_project)
    approved(mini_project)
    (processed / "trend_keywords.csv").write_text("keyword\nchanged\n", encoding="utf-8")
    with pytest.raises(Exception, match="hash|differs|SHA"):
        load_service(mini_project)


def test_real_loader_serves_a_refreshed_extract_with_historical_status(mini_project):
    from job_market_web import create_app
    from job_market_web.build import load_service
    approved(mini_project)
    metrics = build_metrics(mini_project, None, label_state=pv.WRONG_SOURCE)
    served_project(mini_project, metrics_extra=metrics)
    health = create_app(service=load_service(mini_project)).test_client().get("/api/health").get_json()
    assert health["method"] == "hybrid_word_character_tfidf"
    assert health["method_status"].startswith("chosen by the earlier relevance test") and health["build_approval_ref"]


def test_app_prints_the_environment_warning_and_obeys_strict_mode(mini_project, monkeypatch, capsys):
    import job_market_web.build as app_build
    from job_market.environment import EnvironmentMismatchError
    served_project(mini_project)
    approved(mini_project)
    monkeypatch.setattr(app_build, "current_problems", lambda root: ["numpy: pinned 1.26.4, running 0.0"])
    app_build.load_service(mini_project)                                        # normal mode: starts ...
    assert "WARNING: The app environment differs" in capsys.readouterr().out    # ... and says so
    monkeypatch.setenv("JOB_MARKET_STRICT", "1")
    with pytest.raises(EnvironmentMismatchError):
        app_build.load_service(mini_project)
